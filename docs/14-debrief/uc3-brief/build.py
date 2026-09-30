"""Assembles slides/*.png into PermitFlow-v0.4.0-brief.pptx, one full-bleed image per slide."""
import glob, os
from pptx import Presentation
from pptx.util import Emu
here = os.path.dirname(os.path.abspath(__file__))
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
for png in sorted(glob.glob(os.path.join(here, 'slides', '*.png'))):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
prs.save(os.path.join(here, 'PermitFlow-v0.4.0-brief.pptx'))
